#!/usr/bin/env python3
"""Compara as configuracoes de busca (Etapa 6).

    python scripts/eval_retrieval.py                       # so metricas automaticas
    python scripts/eval_retrieval.py --grades devolvida.xlsx

Duas reguas, com forcas diferentes
----------------------------------
  automatica   a aula-fonte de cada pergunta (perguntas_treino.json), mais as
               aulas relacionadas. Gratis e imediata, mas binaria e grosseira:
               acerta a AULA, nao diz se o trecho trazido responde.

  medico       nota 0/1/2 por trecho (avaliacao_trechos.xlsx devolvida). E a
               regua que decide: diz se o trecho ajuda a responder.

Chunk sem nota (abaixo do top-3, que nao entram no pool) herda a nota de um trecho
julgado da mesma aula que o cubra em >= 60%; sem isso fica "nao julgado" e a
cobertura aparece na tabela — metrica com muito buraco nao e conclusao.
"""

from __future__ import annotations

import argparse
import json
import math
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_judging_set import EVAL_DIR, OVERLAP_MIN, POOL_DEPTH, overlap  # noqa: E402
from db import connect  # noqa: E402


def read_grades(path: Path) -> dict[str, int]:
    from openpyxl import load_workbook
    ws = load_workbook(path, data_only=True)["Avaliação"]
    notas = {}
    for trecho, _, nota, *_ in ws.iter_rows(min_row=2, values_only=True):
        if isinstance(trecho, str) and re.fullmatch(r"P\d{2}-T\d{2}", trecho) \
                and nota not in (None, ""):
            notas[trecho] = int(str(nota).strip())
    return notas


def chunk_grades(pool: list[dict], notas: dict[str, int]) -> dict[str, dict[str, int]]:
    """qid -> {chunk_uid: nota}, expandindo cada trecho para todos os chunks que ele representa."""
    out: dict[str, dict[str, int]] = {}
    for p in pool:
        g = out.setdefault(p["qid"], {})
        for t in p["trechos"]:
            if t["excerpt_id"] in notas:
                for u in t["member_chunk_uids"]:
                    g[u] = notas[t["excerpt_id"]]
    return out


def dcg(gains: list[float]) -> float:
    return sum((2 ** g - 1) / math.log2(i + 2) for i, g in enumerate(gains))


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--grades", type=Path, help="planilha devolvida pelo medico")
    ap.add_argument("-k", type=int, default=POOL_DEPTH,
                    help="profundidade; acima do pool julgado a cobertura cai")
    args = ap.parse_args()
    k = args.k

    data = json.loads((EVAL_DIR / "runs.json").read_text("utf-8"))
    runs = data["runs"]
    perguntas = {q["id"]: q for q in
                 json.loads((EVAL_DIR / "perguntas_treino.json").read_text("utf-8"))}

    print(f"Regua automatica (aula-fonte ou relacionada), {len(next(iter(runs.values())))} perguntas\n")
    print(f"{'configuracao':<22} {'acerto@1':>9} {f'acerto@{k}':>9} {'MRR@10':>8}")
    for nome, por_q in runs.items():
        h1 = hk = rr = 0.0
        for qid, hits in por_q.items():
            q = perguntas[qid]
            alvo = {q["lesson_uid"], *q["related_lesson_uids"]}
            pos = next((i for i, h in enumerate(hits, 1) if h["lesson_uid"] in alvo), None)
            h1 += pos == 1
            hk += pos is not None and pos <= k
            rr += 1 / pos if pos else 0
        n = len(por_q)
        print(f"{nome:<22} {h1 / n:>9.0%} {hk / n:>9.0%} {rr / n:>8.3f}")

    if not args.grades:
        print("\n(sem --grades: a regua do medico fica para quando a planilha voltar)")
        return

    pool = json.loads((EVAL_DIR / "judging_pool.json").read_text("utf-8"))
    notas = read_grades(args.grades)
    por_q = chunk_grades(pool, notas)
    print(f"\nRegua do medico: {len(notas)} trechos com nota\n")

    con = connect(read_only=True)
    spans = {r[0]: {"lesson_uid": r[1], "char_start": r[2], "char_end": r[3]}
             for r in con.execute("SELECT chunk_uid, lesson_uid, char_start, char_end "
                                  "FROM chunks").fetchall()}
    con.close()

    def nota(qid: str, uid: str) -> int | None:
        g = por_q.get(qid, {})
        if uid in g:
            return g[uid]
        herdadas = [n for u, n in g.items() if overlap(spans[uid], spans[u]) >= OVERLAP_MIN]
        return max(herdadas) if herdadas else None

    print(f"{'configuracao':<22} {f'nDCG@{k}':>8} {f'P@{k}':>6} {f'sucesso@{k}':>10} {'julgados':>9}")
    for nome, runs_q in runs.items():
        nd = pk = suc = julg = tot = 0.0
        n = 0
        for qid, hits in runs_q.items():
            if qid not in por_q or not por_q[qid]:
                continue
            n += 1
            gs = [nota(qid, h["chunk_uid"]) for h in hits[:k]]
            julg += sum(g is not None for g in gs)
            tot += len(gs)
            gs0 = [g or 0 for g in gs]
            ideal = sorted(por_q[qid].values(), reverse=True)[:k]
            nd += dcg(gs0) / dcg(ideal) if dcg(ideal) else 0
            pk += sum(g >= 1 for g in gs0) / k
            suc += any(g == 2 for g in gs0)
        if n:
            print(f"{nome:<22} {nd / n:>8.3f} {pk / n:>6.0%} {suc / n:>10.0%} "
                  f"{julg / tot:>9.0%}")
    print(f"\nsucesso@{k} = ao menos um trecho nota 2 no top-{k}. "
          "'julgados' baixo = metrica pouco confiavel para aquela linha.")


if __name__ == "__main__":
    main()
