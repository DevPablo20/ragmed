"""Planilha de avaliacao de trechos para o medico (usada por build_judging_set.py).

Escrita para quem nao e de tecnologia: nenhuma palavra como chunk, RAG, rerank
ou configuracao aparece para ele. Tres abas: instrucoes, avaliacao, e um espaco
opcional para as perguntas do proprio medico.
"""

from __future__ import annotations

import math
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

F = "Arial"
AZUL = PatternFill("solid", start_color="1F4E78")
CINZA = PatternFill("solid", start_color="DDEBF7")
AMARELO = PatternFill("solid", start_color="FFF2CC")
VERDE = PatternFill("solid", start_color="E2EFDA")
_fino = Side(style="thin", color="BFBFBF")
BORDA = Border(left=_fino, right=_fino, top=_fino, bottom=_fino)
TOPO = Alignment(wrap_text=True, vertical="top")
TEXTO_COL_LARG = 100

# Exemplo resolvido do tutorial. Pergunta q02, que fica fora do teste; os trechos
# sao reais (encurtados) e foram escolhidos para mostrar uma nota de cada.
EXEMPLO_PERGUNTA = ("Como converter a dose de testosterona em gel para a dose em "
                    "implante hormonal, no homem e na mulher?")
EXEMPLOS = [
    (2, "[...] testosterona dose eu trouxe dose feminina para quem quiser já, né, fazer "
        "a dose de testosterona que está usando em gel. Pra mulher é vezes 25, pra homem "
        "é vezes 5 a 7. [...] Geralmente pra mulher eu uso um miligrama em gel. Quanto que "
        "eu vou botar de implante? 25 [...] se for atleta vezes 7, se for sedentário vezes 5.",
     "Traz exatamente a regra de conversão para os dois sexos. Responde a pergunta."),
    (1, "[...] Pega a dose de gel que tá usando e multiplica [...] eu uso muito para homem, "
        "muito implante de progestional para homem, para evitar aromatização [...] eu "
        "geralmente coloco aí em média 100 miligramas [...] não existe cálculo absoluto, "
        "é cálculo aproximado.",
     "Fala de conversão gel → implante, mas da progesterona, não da testosterona. "
     "Ajuda a entender o raciocínio, mas não responde a pergunta."),
    (0, "[...] Repito, valor total de testosterona é pouco útil no acompanhamento do "
        "tratamento [...] A dose sérica é o resultado do que sobrou. Um paciente que está "
        "precisando de muita testosterona, tu repôs o corpo dele, vai usar tudo [...]",
     "Fala de testosterona, mas sobre interpretar o exame de sangue — não sobre "
     "converter a dose. Mencionar o tema não basta."),
]


def _altura(texto: str, largura: int) -> float:
    """openpyxl nao ajusta altura de linha; estima pelo numero de linhas."""
    linhas = sum(max(1, math.ceil(len(p) / (largura * 1.05))) for p in texto.split("\n"))
    return min(409, 14 * linhas + 6)


def _cel(ws, ref, valor, **fmt):
    c = ws[ref]
    c.value = valor
    c.font = Font(name=F, size=fmt.pop("size", 11), bold=fmt.pop("bold", False),
                  italic=fmt.pop("italic", False), color=fmt.pop("color", None))
    c.alignment = fmt.pop("align", TOPO)
    for k, v in fmt.items():
        setattr(c, k, v)
    return c


def _aba_instrucoes(wb, total: int, n_perguntas: int) -> None:
    ws = wb.active
    ws.title = "Como avaliar"
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 95
    ws.sheet_view.showGridLines = False

    linhas: list[tuple[str, str, dict]] = []   # (col A, col B, estilo)
    def bloco(titulo):
        linhas.append(("", "", {}))
        linhas.append((titulo, "", {"titulo": True}))
    def item(txt, a=""):
        linhas.append((a, txt, {}))

    linhas.append(("Avaliação de trechos de aula", "", {"h1": True}))
    item("Obrigado por ajudar! Leia esta página com calma uma vez (5 minutos) e depois vá "
         "para a aba \"Avaliação\".")

    bloco("Para que serve")
    item("Estou montando um assistente que responde perguntas médicas usando as aulas dos "
         "cursos. Antes de responder, ele procura nas aulas os trechos que tratam do assunto "
         "da pergunta, e a resposta é escrita a partir desses trechos.")
    item("Se ele achar os trechos errados, a resposta sai ruim. Eu preciso saber se ele "
         "está achando os trechos certos — e só um médico consegue julgar isso.")

    bloco("O que você vai fazer")
    item(f"A aba \"Avaliação\" tem {n_perguntas} perguntas. Embaixo de cada pergunta há "
         f"alguns trechos de aula ({total} trechos no total).")
    item("Para cada trecho: leia e escolha uma nota na coluna amarela \"Nota\" "
         "(clique na célula e escolha 0, 1 ou 2 na setinha).")
    item("Não precisa assistir aula nenhuma nem pesquisar nada. Só ler o trecho e "
         "responder: \"este trecho ajuda a responder a pergunta?\"")

    bloco("As notas")
    linhas.append(("2", "RESPONDE — com este trecho dá para responder a pergunta, ou a "
                   "parte principal dela.", {"nota": VERDE}))
    linhas.append(("1", "AJUDA EM PARTE — trata do assunto e traz informação útil, mas "
                   "sozinho não responde (falta a dose, fala de outra situação parecida, "
                   "responde só um pedaço).", {"nota": AMARELO}))
    linhas.append(("0", "NÃO AJUDA — é outro assunto, ou só cita a palavra de passagem.",
                   {"nota": PatternFill("solid", start_color="F8CBAD")}))

    bloco("Regras importantes")
    for t in [
        "Julgue cada trecho sozinho, sem comparar com os outros. Pode haver vários trechos "
        "com nota 2 na mesma pergunta — ou nenhum. As duas coisas são normais.",
        "Julgue se o trecho TRAZ a informação, não se você concorda com ela. Se discordar "
        "do que a aula diz, dê a nota pelo conteúdo e escreva sua opinião na coluna "
        "\"Comentário\". Isso é muito útil para mim.",
        "Os trechos são transcrição automática da fala do professor. Vão aparecer palavras "
        "trocadas, nomes de remédio errados (ex.: \"diasídicos\" no lugar de tiazídicos), "
        "números estranhos e frases soltas. Julgue pelo sentido.",
        "Leia até o fim. Muitas vezes a resposta está no meio ou no final do trecho, depois "
        "de uma conversa com os alunos.",
        "A ordem dos trechos é aleatória de propósito. O primeiro não é o melhor.",
        "Na dúvida entre duas notas, fique com a menor. Se quiser, explique no comentário.",
        "Alguns trechos são parecidos entre si (pedaços vizinhos da mesma aula). Avalie "
        "cada um normalmente.",
        "Se um trecho parecer cortado no fim da célula, clique nele: o texto completo "
        "aparece na barra de cima (ou dê dois cliques na célula).",
    ]:
        item(t, "•")

    bloco("Tempo")
    item(f"São {total} trechos, cerca de 1 a 1,5 minuto cada — por volta de 1 hora no "
         "total. Não precisa fazer de uma vez: salve o arquivo e continue depois de onde "
         "parou. A aba \"Suas perguntas\" é opcional, só se sobrar tempo.")
    linhas.append(("Progresso", "", {"progresso": True}))

    bloco("Exemplo resolvido")
    item(f"Pergunta: {EXEMPLO_PERGUNTA}", "")
    for nota, trecho, porque in EXEMPLOS:
        linhas.append((str(nota), f"Trecho: \"{trecho}\"", {"exemplo": True}))
        item(f"Por quê: {porque}")

    bloco("Opcional: suas perguntas")
    item("Na aba \"Suas perguntas\", escreva até 10 perguntas que você faria no dia a dia "
         "para um assistente desses. Elas entram na próxima rodada de testes e são as mais "
         "valiosas, porque são reais.")

    bloco("Quando terminar")
    item("Salve o arquivo e me mande de volta. Pode mandar mesmo que não tenha terminado "
         "tudo — o que estiver avaliado já serve.")

    for i, (a, b, st) in enumerate(linhas, 1):
        if st.get("h1"):
            _cel(ws, f"A{i}", a, size=16, bold=True, align=Alignment(vertical="center"))
            ws.row_dimensions[i].height = 28
            continue
        if st.get("titulo"):
            _cel(ws, f"A{i}", a, size=12, bold=True, color="1F4E78")
            continue
        if st.get("progresso"):
            _cel(ws, f"A{i}", a, bold=True)
            c = _cel(ws, f"B{i}",
                     f"=\"Você já avaliou \"&(COUNTA('Avaliação'!C:C)-1)&\" de {total} "
                     f"trechos.\"", bold=True, color="375623")
            c.fill = VERDE
            continue
        ca = _cel(ws, f"A{i}", a, bold=bool(st.get("nota") or st.get("exemplo")),
                  size=12 if st.get("nota") or st.get("exemplo") else 11,
                  align=Alignment(horizontal="center", vertical="top")
                  if (st.get("nota") or st.get("exemplo")) else TOPO)
        cb = _cel(ws, f"B{i}", b, italic=bool(st.get("exemplo")))
        if st.get("nota"):
            ca.fill = st["nota"]
        ws.row_dimensions[i].height = _altura(b, 95) if b else 15


def _aba_avaliacao(wb, pool: list[dict]) -> None:
    ws = wb.create_sheet("Avaliação")
    for col, w in zip("ABCD", [11, TEXTO_COL_LARG, 9, 36]):
        ws.column_dimensions[col].width = w
    for col, h in zip("ABCD", ["Trecho", "Texto do trecho", "Nota", "Comentário (opcional)"]):
        c = _cel(ws, f"{col}1", h, bold=True, color="FFFFFF",
                 align=Alignment(vertical="center", wrap_text=True))
        c.fill, c.border = AZUL, BORDA
    ws.freeze_panes = "A2"

    dv = DataValidation(type="list", formula1='"0,1,2"', allow_blank=True,
                        showErrorMessage=True, errorTitle="Nota inválida",
                        error="Use 0, 1 ou 2.")
    ws.add_data_validation(dv)

    r = 2
    for p in pool:
        ws.merge_cells(f"A{r}:D{r}")
        c = _cel(ws, f"A{r}", f"PERGUNTA {p['numero']} de {len(pool)}:  {p['pergunta']}",
                 size=12, bold=True, align=Alignment(wrap_text=True, vertical="center"))
        c.fill = CINZA
        ws.row_dimensions[r].height = 36
        r += 1
        for t in p["trechos"]:
            _cel(ws, f"A{r}", t["excerpt_id"], size=9, color="7F7F7F").border = BORDA
            _cel(ws, f"B{r}", t["text"].strip(), size=10).border = BORDA
            for col in "CD":
                c = _cel(ws, f"{col}{r}", None,
                         align=Alignment(horizontal="center", vertical="top")
                         if col == "C" else TOPO)
                c.fill, c.border = AMARELO, BORDA
            dv.add(f"C{r}")
            ws.row_dimensions[r].height = _altura(t["text"].strip(), TEXTO_COL_LARG)
            r += 1
        r += 1  # linha em branco entre perguntas


def _aba_perguntas(wb) -> None:
    ws = wb.create_sheet("Suas perguntas")
    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 80
    ws.column_dimensions["C"].width = 45
    _cel(ws, "A1", "Opcional: perguntas que você faria no dia a dia (até 10).", bold=True,
         size=12)
    _cel(ws, "A2", "Escreva do jeito que você perguntaria a um colega. Não precisa saber em "
         "qual aula está a resposta.", italic=True)
    for col, h in zip("ABC", ["#", "Pergunta", "Situação em que você perguntaria (opcional)"]):
        c = _cel(ws, f"{col}4", h, bold=True, color="FFFFFF")
        c.fill, c.border = AZUL, BORDA
    for i in range(1, 11):
        r = 4 + i
        _cel(ws, f"A{r}", i).border = BORDA
        for col in "BC":
            c = _cel(ws, f"{col}{r}", None)
            c.fill, c.border = AMARELO, BORDA
        ws.row_dimensions[r].height = 32


def write_sheet(pool: list[dict], path: Path) -> None:
    total = sum(len(p["trechos"]) for p in pool)
    wb = Workbook()
    _aba_instrucoes(wb, total, len(pool))
    _aba_avaliacao(wb, pool)
    _aba_perguntas(wb)
    wb.save(path)
